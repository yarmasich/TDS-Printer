from datetime import datetime
from pathlib import Path

import pytest
from fastapi import HTTPException, UploadFile
from sqlalchemy import event
from sqlmodel import Session, SQLModel, create_engine, select
from openpyxl import Workbook

from app.models import CartItem, DataHall, Discipline, Import, Label, Project
from app.importer import InvalidWorkbookError, scan_labels_dir
from app.api.projects import delete_discipline, DisciplinePatch, DisciplineIn, create_discipline
from app.api.imports import delete_import

@pytest.fixture
def import_engine():
    engine = create_engine('sqlite://')
    @event.listens_for(engine, 'connect')
    def configure(connection, _):
        connection.execute('PRAGMA foreign_keys=ON')
    SQLModel.metadata.create_all(engine)
    yield engine
    engine.dispose()

def workbook(path, text='CBL#1.1'):
    wb = Workbook()
    wb.active.append([text, text])
    wb.save(path)
    wb.close()

def tree(tmp_path):
    directory = tmp_path / 'P' / 'H' / 'D'
    directory.mkdir(parents=True)
    return directory

def test_scan_commits_empty_directories(import_engine, tmp_path):
    tree(tmp_path)
    with Session(import_engine) as session:
        assert scan_labels_dir(session, tmp_path).disciplines_created == 1
    with Session(import_engine) as session:
        assert len(session.exec(select(Discipline)).all()) == 1

def test_failed_late_scan_restores_all_previous_data(import_engine, tmp_path):
    directory = tree(tmp_path)
    workbook(directory / 'a.xlsx', 'original')
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
    workbook(directory / 'a.xlsx', 'replacement')
    (directory / 'z.xlsx').write_bytes(b'broken workbook')
    with Session(import_engine) as session:
        with pytest.raises(InvalidWorkbookError):
            scan_labels_dir(session, tmp_path, wipe=True)
    with Session(import_engine) as session:
        assert [label.left_text for label in session.exec(select(Label)).all()] == ['original']
        assert len(session.exec(select(Import)).all()) == 1

@pytest.mark.parametrize('operation', ['import', 'discipline', 'wipe'])
def test_delete_removes_queued_cart_dependencies(import_engine, tmp_path, operation):
    directory = tree(tmp_path)
    workbook(directory / 'a.xlsx')
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        label = session.exec(select(Label)).one()
        session.add(CartItem(session_id='s', label_id=label.id))
        session.commit()
        if operation == 'import':
            delete_import(label.import_id, session)
        elif operation == 'discipline':
            delete_discipline(label.discipline_id, session)
        else:
            scan_labels_dir(session, tmp_path, wipe=True)
        assert session.exec(select(CartItem)).all() == []
        assert len(session.exec(select(Import)).all()) == (1 if operation == 'wipe' else 0)

@pytest.mark.parametrize('status', ['printing', 'uncertain'])
@pytest.mark.parametrize('operation', ['import', 'discipline', 'wipe'])
def test_deletion_rejects_active_jobs(import_engine, tmp_path, status, operation):
    directory = tree(tmp_path)
    workbook(directory / 'a.xlsx')
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        label = session.exec(select(Label)).one()
        item = CartItem(session_id='s', label_id=label.id, status=status)
        session.add(item)
        session.commit()
        with pytest.raises(HTTPException) as exc:
            if operation == 'import': delete_import(label.import_id, session)
            elif operation == 'discipline': delete_discipline(label.discipline_id, session)
            else: scan_labels_dir(session, tmp_path, wipe=True)
        assert exc.value.status_code == 409

@pytest.mark.parametrize('field', ['name', 'color', 'bundle_mode'])
def test_patch_rejects_required_null(field):
    with pytest.raises(ValueError):
        DisciplinePatch(**{field: None})

def test_create_missing_template_returns_4xx(import_engine):
    with Session(import_engine) as session:
        p = Project(name='P'); session.add(p); session.commit()
        h = DataHall(project_id=p.id, name='H'); session.add(h); session.commit()
        with pytest.raises(HTTPException) as exc:
            create_discipline(DisciplineIn(data_hall_id=h.id, name='D', template_id=999), session)
        assert 400 <= exc.value.status_code < 500


def test_upload_server_owned_filename_and_invalid_workbook(import_engine, tmp_path):
    import asyncio
    from io import BytesIO
    from app.api.import_ import upload
    directory = tree(tmp_path)
    workbook(directory / 'a.xlsx')
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        discipline = session.exec(select(Discipline)).one()
        victim = tmp_path / 'victim.xlsx'
        victim.write_bytes(b'keep me')
        data = (directory / 'a.xlsx').read_bytes()
        result = asyncio.run(upload(discipline.id, UploadFile(filename=str(victim), file=BytesIO(data)), session))
        assert result['rows'] == 1
        assert victim.read_bytes() == b'keep me'
        with pytest.raises(HTTPException) as exc:
            asyncio.run(upload(discipline.id, UploadFile(filename='broken.xlsx', file=BytesIO(b'bad')), session))
        assert exc.value.status_code == 400
        assert len(session.exec(select(Import)).all()) == 2


def test_initial_import_migration_preserves_populated_cart(monkeypatch):
    from alembic.config import Config
    from alembic import command
    from sqlalchemy import inspect, text
    import app.db as db
    engine = create_engine('sqlite://')
    @event.listens_for(engine, 'connect')
    def configure(connection, _): connection.execute('PRAGMA foreign_keys=ON')
    monkeypatch.setattr(db, 'engine', engine)
    root = Path(__file__).resolve().parents[1]
    config = Config(str(root / 'alembic.ini'))
    config.set_main_option('script_location', str(root / 'alembic'))
    command.upgrade(config, 'eb677bd2f0fc')
    with engine.begin() as connection:
        for model in [Project(id=1, name='P'), DataHall(id=1, project_id=1, name='H'),
                      Discipline(id=1, data_hall_id=1, name='D'), Label(id=1, discipline_id=1),
                      CartItem(id=1, session_id='s', label_id=1)]:
            columns = {c['name'] for c in inspect(connection).get_columns(model.__tablename__)}
            data = {key: value.isoformat() if isinstance(value, datetime) else value
                    for key, value in model.model_dump().items() if key in columns}
            connection.execute(text('INSERT INTO ' + model.__tablename__ + ' (' + ','.join(data) + ') VALUES (' + ','.join(':' + key for key in data) + ')'), data)
    command.upgrade(config, '232e2ff8080e')
    with engine.connect() as connection:
        assert connection.exec_driver_sql('PRAGMA foreign_keys').scalar() == 1
        assert connection.exec_driver_sql('PRAGMA foreign_key_check').all() == []
        assert connection.exec_driver_sql('SELECT label_id FROM cartitem').scalar() == 1
    command.upgrade(config, 'head')
    command.check(config)
    with engine.connect() as connection:
        assert connection.exec_driver_sql('PRAGMA foreign_key_check').all() == []
        assert connection.exec_driver_sql('SELECT label_id, status, claim_token FROM cartitem').one() == (1, 'queued', '')
    engine.dispose()


def test_duplicate_project_name_is_conflict(import_engine):
    from app.api.projects import ProjectIn, create_project
    with Session(import_engine) as session:
        create_project(ProjectIn(name='P'), session)
        with pytest.raises(HTTPException) as exc:
            create_project(ProjectIn(name=' P '), session)
        assert exc.value.status_code == 409
        assert len(session.exec(select(Project)).all()) == 1


@pytest.mark.parametrize('level', ['hall', 'project'])
def test_parent_deletion_handles_all_dependencies(import_engine, tmp_path, level):
    from app.api.projects import delete_hall, delete_project
    directory = tree(tmp_path)
    workbook(directory / 'a.xlsx')
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        label = session.exec(select(Label)).one()
        session.add(CartItem(session_id='s', label_id=label.id))
        session.commit()
        if level == 'hall': delete_hall(session.exec(select(DataHall.id)).one(), session)
        else: delete_project(session.exec(select(Project.id)).one(), session)
        for model in [CartItem, Label, Import, Discipline, DataHall]:
            assert session.exec(select(model)).all() == []


def test_parent_active_job_rolls_back_already_deleted_sibling(import_engine, tmp_path):
    from app.api.projects import delete_project
    first = tree(tmp_path)
    second = first.parent / 'Z'; second.mkdir()
    workbook(first / 'a.xlsx'); workbook(second / 'a.xlsx')
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        labels = session.exec(select(Label).order_by(Label.id)).all()
        session.add(CartItem(session_id='s', label_id=labels[-1].id, status='uncertain'))
        session.commit()
        project_id = session.exec(select(Project.id)).one()
        with pytest.raises(HTTPException) as exc:
            delete_project(project_id, session)
        assert exc.value.status_code == 409
        assert len(session.exec(select(Label)).all()) == 2
        assert len(session.exec(select(Discipline)).all()) == 2


def test_workbook_helper_does_not_commit_callers_transaction(import_engine, tmp_path):
    from app.importer import import_workbook_into_discipline
    directory = tree(tmp_path)
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        discipline = session.exec(select(Discipline)).one()
        workbook(directory / 'a.xlsx')
        import_workbook_into_discipline(session, directory / 'a.xlsx', discipline)
        session.rollback()
    with Session(import_engine) as session:
        assert session.exec(select(Import)).all() == []
        assert session.exec(select(Label)).all() == []


@pytest.mark.parametrize('setting, value, detail', [
    ('TDS_IMPORT_MAX_FILE_BYTES', '32', 'file size'),
    ('TDS_IMPORT_MAX_UNCOMPRESSED_BYTES', '1024', 'uncompressed'),
    ('TDS_IMPORT_MAX_ARCHIVE_ENTRIES', '2', 'archive entries'),
    ('TDS_IMPORT_MAX_SHEETS', '1', 'sheets'),
    ('TDS_IMPORT_MAX_ROWS', '2', 'rows'),
])
def test_upload_quota_rejects_and_rolls_back(import_engine, tmp_path, monkeypatch, setting, value, detail):
    import asyncio
    from io import BytesIO
    from app.api.import_ import upload
    directory = tree(tmp_path)
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        discipline = session.exec(select(Discipline)).one()
        wb = Workbook()
        wb.active.append(['one', 'one'])
        wb.active.append([None, None])
        wb.active.append(['three', 'three'])
        wb.create_sheet('second')
        path = directory / 'a.xlsx'
        wb.save(path); wb.close()
        monkeypatch.setenv(setting, value)
        with pytest.raises(HTTPException) as exc:
            asyncio.run(upload(discipline.id, UploadFile(filename='a.xlsx', file=BytesIO(path.read_bytes())), session))
        assert exc.value.status_code == 413
        assert detail in exc.value.detail.lower()
        assert session.exec(select(Import)).all() == []
        assert session.exec(select(Label)).all() == []


def test_scan_quota_rolls_back_wipe(import_engine, tmp_path, monkeypatch):
    from app.api.import_ import scan
    directory = tree(tmp_path)
    workbook(directory / 'a.xlsx', 'original')
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        monkeypatch.setenv('TDS_IMPORT_MAX_FILE_BYTES', '32')
        with pytest.raises(HTTPException) as exc:
            scan(str(tmp_path), True, session)
        assert exc.value.status_code == 413
        assert [label.left_text for label in session.exec(select(Label)).all()] == ['original']


def test_rows_quota_is_workbook_wide_including_blanks(import_engine, tmp_path, monkeypatch):
    import asyncio
    from io import BytesIO
    from app.api.import_ import upload
    directory = tree(tmp_path)
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        discipline = session.exec(select(Discipline)).one()
        wb = Workbook()
        wb.active.append(['one'])
        wb.create_sheet('second').append(['two'])
        path = directory / 'a.xlsx'; wb.save(path); wb.close()
        monkeypatch.setenv('TDS_IMPORT_MAX_ROWS', '1')
        with pytest.raises(HTTPException) as exc:
            asyncio.run(upload(discipline.id, UploadFile(filename='a.xlsx', file=BytesIO(path.read_bytes())), session))
        assert exc.value.status_code == 413
        assert session.exec(select(Label)).all() == []


def test_upload_copy_stops_at_limit_plus_one(monkeypatch):
    from io import BytesIO
    from app.import_limits import ImportLimitError, ImportLimits, copy_upload
    monkeypatch.setenv('TDS_IMPORT_MAX_FILE_BYTES', '10')
    source = BytesIO(b'x' * 1000)
    target = BytesIO()
    with pytest.raises(ImportLimitError):
        copy_upload(source, target, ImportLimits.from_env())
    assert source.tell() == 11
    assert len(target.getvalue()) <= 10


def test_import_ignores_false_dimensions_and_extra_columns(import_engine, tmp_path, monkeypatch):
    from zipfile import ZipFile
    from app.importer import import_workbook_into_discipline
    directory = tree(tmp_path)
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        discipline = session.exec(select(Discipline)).one()
        source = directory / 'source.xlsx'
        wb = Workbook(); wb.active.append(['left', 'right', 'ignored'])
        wb.save(source); wb.close()
        path = directory / 'edited.xlsx'
        with ZipFile(source) as archive, ZipFile(path, 'w') as edited:
            for entry in archive.infolist():
                content = archive.read(entry)
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    content = content.replace(b'ref="A1:C1"', b'ref="A1:XFD1048576"')
                edited.writestr(entry, content)
        monkeypatch.setenv('TDS_IMPORT_MAX_ROWS', '1')
        stats = import_workbook_into_discipline(session, path, discipline)
        assert stats.rows == 1
        label = session.exec(select(Label)).one()
        assert (label.left_text, label.right_text) == ('left', 'right')


def test_corrupt_compressed_member_returns_400(import_engine, tmp_path):
    import asyncio
    from io import BytesIO
    from zipfile import ZipFile, ZIP_DEFLATED
    from app.api.import_ import upload
    directory = tree(tmp_path)
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        discipline = session.exec(select(Discipline)).one()
        buffer = BytesIO()
        with ZipFile(buffer, 'w', compression=ZIP_DEFLATED) as archive:
            archive.writestr('bad.xml', b'hello' * 100)
        content = bytearray(buffer.getvalue())
        # Invalid DEFLATE block type in the first member, past its local header.
        content[30 + len('bad.xml')] = 0xff
        with pytest.raises(HTTPException) as exc:
            asyncio.run(upload(discipline.id, UploadFile(filename='bad.xlsx', file=BytesIO(content)), session))
        assert exc.value.status_code == 400
        assert session.exec(select(Import)).all() == []


@pytest.mark.parametrize('kind', ['encrypted', 'unsupported'])
def test_unsupported_zip_member_returns_400(import_engine, tmp_path, kind):
    import asyncio
    from io import BytesIO
    from zipfile import ZipFile
    from app.api.import_ import upload
    directory = tree(tmp_path)
    with Session(import_engine) as session:
        scan_labels_dir(session, tmp_path)
        discipline = session.exec(select(Discipline)).one()
        buffer = BytesIO()
        with ZipFile(buffer, 'w') as archive:
            archive.writestr('bad.xml', b'hello')
        content = bytearray(buffer.getvalue())
        central = content.index(b'PK\x01\x02')
        if kind == 'encrypted':
            content[6] |= 1
            content[central + 8] |= 1
        else:
            content[8:10] = (99).to_bytes(2, 'little')
            content[central + 10:central + 12] = (99).to_bytes(2, 'little')
        with pytest.raises(HTTPException) as exc:
            asyncio.run(upload(discipline.id, UploadFile(filename='bad.xlsx', file=BytesIO(content)), session))
        assert exc.value.status_code == 400
        assert session.exec(select(Import)).all() == []
